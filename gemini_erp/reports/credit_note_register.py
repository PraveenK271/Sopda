"""Credit Notes Received Register — Excel and PDF (patch v1.1.0, P4).

For reconciling recorded supplier credit notes against GSTR-2B / the IMS
dashboard. Renders rows exactly as CreditNoteService.list_credit_notes()
returns them; the only arithmetic here is the column totals on the last row.
"""

import logging
from datetime import date as date_type
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

logger = logging.getLogger(__name__)

# (header, row key, is_amount)
COLUMNS = [
    ("Date", "credit_note_date", False),
    ("Supplier", "supplier_name", False),
    ("GSTIN", "supplier_gstin", False),
    ("Credit Note No", "credit_note_no", False),
    ("Type", "note_type", False),
    ("Reason", "reason", False),
    ("Against Invoice", "linked_invoice_no", False),
    ("Taxable", "taxable_value", True),
    ("CGST", "cgst", True),
    ("SGST", "sgst", True),
    ("IGST", "igst", True),
    ("Total", "total_amount", True),
    ("ITC Reversed", "itc_reversal_amount", True),
    ("IMS Status", "ims_status", False),
    ("Declare By", "declaration_deadline", False),
    ("Flag", "flag", False),
]


def _cell(row: dict, key: str):
    if key == "flag":
        return {"PASSED": "Cut-off passed", "APPROACHING": "Cut-off near"}.get(
            row.get("deadline_status") if row.get("needs_attention") else None, ""
        )
    value = row.get(key)
    if isinstance(value, date_type):
        return value.strftime("%d-%m-%Y")
    return "" if value is None else value


def _totals(rows: list[dict]) -> dict:
    return {key: round(sum(r[key] for r in rows), 2) for _h, key, is_amount in COLUMNS if is_amount}


def export_excel(rows: list[dict], path: str, title: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Credit Notes Received"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([header for header, _k, _a in COLUMNS])
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in rows:
        ws.append([_cell(row, key) for _h, key, _a in COLUMNS])
    totals = _totals(rows)
    ws.append(["Total"] + [totals.get(key, "") for _h, key, _a in COLUMNS[1:]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for index, (_h, key, is_amount) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = 14 if is_amount else 18
        if is_amount:
            for (cell,) in ws.iter_rows(min_row=3, min_col=index, max_col=index):
                cell.number_format = "#,##0.00"
    wb.save(path)
    logger.info("Credit notes register (%d rows) exported to %s", len(rows), path)
    return path


def export_pdf(rows: list[dict], path: str, title: str) -> str:
    styles = getSampleStyleSheet()
    small = styles["Normal"].clone("Small", fontSize=7, leading=8)
    data = [[header for header, _k, _a in COLUMNS]]
    for row in rows:
        data.append([
            f"{_cell(row, key):,.2f}" if is_amount else Paragraph(escape(str(_cell(row, key))), small)
            for _h, key, is_amount in COLUMNS
        ])
    totals = _totals(rows)
    data.append(["Total"] + [
        f"{totals[key]:,.2f}" if is_amount else "" for _h, key, is_amount in COLUMNS[1:]
    ])

    widths = [16, 30, 26, 20, 16, 20, 18, 16, 13, 13, 13, 16, 14, 14, 16, 16]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (7, 1), (12, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    doc = SimpleDocTemplate(
        path, pagesize=landscape(A4),
        leftMargin=8 * mm, rightMargin=8 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    doc.build([Paragraph(escape(title), styles["Heading3"]), Spacer(1, 3 * mm), table])
    logger.info("Credit notes register (%d rows) exported to %s", len(rows), path)
    return path
